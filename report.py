"""Model report.

1. collect()  - the host calls the read-only tools itself, in a fixed order,
                and stores everything in one JSON (no model involved).
2. write_text() - the local model writes the prose sections from that JSON only.
3. build_docx() - tables come straight from the JSON; the prose is inserted.

Numbers in the tables therefore never pass through the model.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Iterator

import ollama
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import mechanical
from mcp_host import McpHub, parse_json


# ---------------------------------------------------------------------- 1. collect
def _tool(hub: McpHub, name: str, args: dict | None = None):
    data = parse_json(hub.call(name, args or {}, timeout_s=180))
    return data if isinstance(data, dict) else {}


def collect(hub: McpHub) -> dict:
    info = _tool(hub, "get_model_info").get("model_info", {})
    analyses = mechanical.results_overview(hub)
    if not analyses:  # fall back to the plain tools if the script path failed
        analyses = [
            {"index": i, "name": a.get("name"), "type": a.get("type"), "results": []}
            for i, a in enumerate(info.get("analyses", []))
        ]
        for a in analyses:
            a["status"] = _tool(hub, "get_solve_status", {"analysis_index": a["index"]}).get("status")

    for a in analyses:
        idx = a["index"]
        a["boundary_conditions"] = _tool(
            hub, "list_boundary_conditions", {"analysis_index": idx}
        ).get("boundary_conditions", [])
        if "modal" in str(a.get("type", "")).lower():
            a["frequencies"] = _tool(
                hub, "get_modal_frequencies", {"analysis_index": idx}
            ).get("frequencies", [])

    return {
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "connections": mechanical.connections(hub),
        "bodies": info.get("bodies", []),
        "materials": _tool(hub, "list_materials").get("materials", []),
        "named_selections": _tool(hub, "list_named_selections").get("named_selections", []),
        "point_masses": _tool(hub, "list_point_masses").get("point_masses", []),
        "mesh": _tool(hub, "get_mesh_statistics").get("mesh_statistics", {}),
        "analyses": analyses,
    }


# ---------------------------------------------------------------------- 2. text
_TEXT_SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "model_description": {"type": "string"},
        "results_discussion": {"type": "string"},
        "notes": {"type": "string"},
    },
    "required": ["summary", "model_description", "results_discussion", "notes"],
}

_TEXT_PROMPT = """Write the prose sections of an engineering report about an ANSYS Mechanical model.
The data below was read from the model. It is the only source you may use.

Rules:
- English, technical, concise, third person.
- Use only names and numbers that appear in the data, with their units. Do not invent loads,
  limits, safety factors or conclusions about acceptability; no acceptance criteria were given.
- If a result has no value or an analysis is not solved, say so plainly.
- summary: 3-5 sentences. model_description: bodies, materials, mesh, supports and loads.
- results_discussion: what the evaluated results and frequencies show, per analysis.
- notes: missing data, unsolved analyses, or items the engineer should check. Plain sentences.

DATA:
{data}"""


def write_text(data: dict, cfg: dict) -> dict:
    o = cfg["ollama"]
    client = ollama.Client(host=o["url"], timeout=o.get("timeout_s", 600))
    kwargs = {}
    if "think" in o:
        kwargs["think"] = o["think"]
    resp = client.chat(
        model=o["model"],
        messages=[{"role": "user", "content": _TEXT_PROMPT.format(data=json.dumps(data, indent=1))}],
        format=_TEXT_SCHEMA,
        options={"num_ctx": o.get("num_ctx", 16384), "temperature": 0.1},
        **kwargs,
    )
    text = parse_json(resp.message.content) or {}
    return {k: str(text.get(k, "")).strip() for k in _TEXT_SCHEMA["properties"]}


# ---------------------------------------------------------------------- 3. docx
_INK = RGBColor(0x1B, 0x24, 0x30)


def _table(doc: Document, headers: list[str], rows: list[list]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "-" if value in (None, "") else str(value)
    doc.add_paragraph()


def _paragraphs(doc: Document, text: str) -> None:
    for block in [b.strip() for b in text.split("\n") if b.strip()]:
        doc.add_paragraph(block)


def _pick(images: list[dict], category: str, analysis: int | None = None) -> list[dict]:
    out = [i for i in images if i.get("category") == category and i.get("path")]
    if analysis is not None:
        out = [i for i in out if i.get("analysis") == analysis]
    return out


def _caption(doc: Document, text: str) -> None:
    run = doc.add_paragraph(text).runs[0]
    run.italic = True
    run.font.size = Pt(9)


def _gallery(doc: Document, images: list[dict], cfg: dict, columns: int = 3) -> None:
    """Few images: one per page width. Many: a grid of thumbnails with names."""
    if not images:
        return
    big = cfg.get("report", {}).get("full_size_up_to", 2)
    if len(images) <= big:
        for img in images:
            try:
                doc.add_picture(img["path"], width=Inches(cfg.get("report", {}).get("image_width_in", 5.8)))
                _caption(doc, f"{img.get('name')} ({img.get('type', '')})".replace(" ()", ""))
            except Exception:  # noqa: BLE001 - an unreadable file must not kill the report
                pass
        return

    thumb = Inches(cfg.get("report", {}).get("thumb_width_in", 1.7))
    rows = (len(images) + columns - 1) // columns
    table = doc.add_table(rows=rows * 2, cols=columns)
    table.style = "Table Grid"
    for n, img in enumerate(images):
        r, c = divmod(n, columns)
        try:
            table.rows[r * 2].cells[c].paragraphs[0].add_run().add_picture(img["path"], width=thumb)
        except Exception:  # noqa: BLE001
            table.rows[r * 2].cells[c].text = "(image failed)"
        label = table.rows[r * 2 + 1].cells[c].paragraphs[0]
        run = label.add_run(str(img.get("name", "")))
        run.font.size = Pt(8)
        run.bold = True
    doc.add_paragraph()


def build_docx(data: dict, text: dict, cfg: dict, out_path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading("ANSYS Mechanical model report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = doc.add_paragraph()
    meta.add_run(f"Generated {data['generated']} by {cfg['app']['title']} ").font.color.rgb = _INK
    meta.add_run(f"({cfg['ollama']['model']}, local)").italic = True

    doc.add_heading("1. Summary", 1)
    _paragraphs(doc, text["summary"] or "No summary was generated.")

    doc.add_heading("2. Model", 1)
    _paragraphs(doc, text["model_description"])
    doc.add_heading("2.1 Bodies", 2)
    _table(doc, ["Body", "Material"], [[b.get("name"), b.get("material")] for b in data["bodies"]])
    _gallery(doc, _pick(data.get("images", []), "geometry"), cfg)
    if data["materials"]:
        doc.add_heading("2.2 Materials in Engineering Data", 2)
        _table(doc, ["Material"], [[m] for m in data["materials"]])

    doc.add_heading("3. Connections", 1)
    conns = data.get("connections", [])
    if conns:
        _table(doc, ["Group", "Connection", "Type", "Behavior", "Formulation"],
               [[c.get("group"), c.get("name"), c.get("ContactType") or c.get("type"),
                 c.get("Behavior"), c.get("FormulationType")] for c in conns])
        _gallery(doc, _pick(data.get("images", []), "connections"), cfg)
    else:
        doc.add_paragraph("No contacts, joints or springs were found in the model.")

    doc.add_heading("4. Mesh", 1)
    _table(doc, ["Nodes", "Elements"],
           [[data["mesh"].get("nodes"), data["mesh"].get("elements")]])
    _gallery(doc, _pick(data.get("images", []), "mesh"), cfg)

    if data["named_selections"]:
        doc.add_heading("5. Named selections", 1)
        _table(doc, ["Named selection", "Entities"],
               [[n.get("name"), n.get("entity_count")] for n in data["named_selections"]])
        _gallery(doc, _pick(data.get("images", []), "named_selections"), cfg)
    if data["point_masses"]:
        doc.add_heading("5.1 Point masses", 2)
        _table(doc, ["Name", "Mass [kg]", "CG [mm]", "Scoped to"],
               [[p.get("name"), p.get("mass_kg"), p.get("cg_mm"), p.get("named_selection")]
                for p in data["point_masses"]])

    doc.add_heading("6. Analyses", 1)
    for a in data["analyses"]:
        idx = a["index"]
        doc.add_heading(f"6.{idx + 1} {a.get('name')} ({a.get('type')})", 2)
        doc.add_paragraph(f"Solution status: {a.get('status', '-')}")

        bcs = [b for b in a.get("boundary_conditions", [])
               if b.get("type") not in ("ANSYSAnalysisSettings", "AnalysisSettings", "Solution")]
        doc.add_heading("Supports and loads", 3)
        if bcs:
            _table(doc, ["Object", "Type"], [[b.get("name"), b.get("type")] for b in bcs])
            _gallery(doc, _pick(data.get("images", []), "boundary_conditions", idx), cfg)
        else:
            doc.add_paragraph("No supports or loads are defined.")

        results = [r for r in a.get("results", []) if r.get("type") != "SolutionInformation"]
        doc.add_heading("Results", 3)
        if results:
            _table(doc, ["Result", "Type", "Minimum", "Maximum", "State"],
                   [[r.get("name"), r.get("type"), r.get("minimum"), r.get("maximum"), r.get("state")]
                    for r in results])
            _gallery(doc, _pick(data.get("images", []), "results", idx), cfg)
        else:
            doc.add_paragraph("No result objects are defined.")
        if a.get("frequencies"):
            _table(doc, ["Mode", "Frequency [Hz]"],
                   [[f.get("mode"), f.get("frequency_hz")] for f in a["frequencies"]])

    doc.add_heading("7. Discussion", 1)
    _paragraphs(doc, text["results_discussion"])
    doc.add_heading("8. Notes", 1)
    _paragraphs(doc, text["notes"])

    note = doc.add_paragraph()
    run = note.add_run(
        "Tables and images are taken directly from the Mechanical model. The summary, the model "
        "description, the discussion and the notes are written by a local language model from the "
        "same data and should be reviewed."
    )
    run.italic = True
    run.font.size = Pt(8.5)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------- pipeline
def run_report(hub: McpHub, cfg: dict, out_dir: Path) -> Iterator[dict]:
    """Generator of progress events; last event carries the file names."""
    state = mechanical.status(hub)
    if not state["connected"]:
        yield {"type": "error", "message": "Mechanical is not connected. Press Connect first."}
        return
    yield {"type": "step", "text": "Reading the model"}
    data = collect(hub)
    if cfg.get("report", {}).get("images", True):
        yield {"type": "step", "text": "Exporting images from Mechanical (one per tree object)"}
        data["images"] = mechanical.export_images(
            hub, out_dir / "images", cfg.get("report", {}).get("max_images", 80))
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = out_dir / f"model_report_{stamp}.json"
    out_dir.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), "utf-8")

    yield {"type": "step", "text": "Writing the text sections"}
    try:
        text = write_text(data, cfg)
    except Exception as exc:  # noqa: BLE001 - still produce the tables
        text = {k: "" for k in _TEXT_SCHEMA["properties"]}
        text["notes"] = f"The text sections could not be generated: {exc}"

    yield {"type": "step", "text": "Building the Word file"}
    docx_path = build_docx(data, text, cfg, out_dir / f"model_report_{stamp}.docx")
    yield {"type": "report", "docx": docx_path.name, "json": json_path.name}
